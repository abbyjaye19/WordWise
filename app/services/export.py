"""Export a reading the way the student set it up: their highlights, font,
text size, spacing, and background.

Word indexes match assets/studio.js: words are runs of non-space characters,
counted in order across every paragraph of the document.
"""

from __future__ import annotations

import html
import io
import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

WORD_RE = re.compile(r"\S+")

FONTS = {"atkinson": "Atkinson Hyperlegible", "lexend": "Lexend", "verdana": "Verdana"}
CSS_FONTS = {
    "atkinson": "'Atkinson Hyperlegible', Verdana, sans-serif",
    "lexend": "'Lexend', Verdana, sans-serif",
    "verdana": "Verdana, Tahoma, sans-serif",
}
SIZES_PT = {"s": 12, "m": 14, "l": 16, "xl": 18}
SIZES_PX = {"s": 16, "m": 18, "l": 21, "xl": 24}
SPACING = {  # line height, letter spacing (em), word spacing (em)
    "normal": (1.6, 0.0, 0.0),
    "wide": (1.9, 0.03, 0.08),
    "xwide": (2.2, 0.06, 0.16),
}
THEMES = {  # background, text
    "cream": ("#FBF8F1", "#0F1F3D"),
    "green": ("#EEF6EB", "#0F1F3D"),
    "blue": ("#EEF2F8", "#0F1F3D"),
    "dark": ("#13213D", "#F4EEE1"),
}
HL_CSS = {"gold": "#F1DFB5", "green": "#CFE3D3", "blue": "#D3DDF0"}
HL_DOCX = {"gold": WD_COLOR_INDEX.YELLOW, "green": WD_COLOR_INDEX.BRIGHT_GREEN, "blue": WD_COLOR_INDEX.TURQUOISE}


def _prefs(prefs: dict | None) -> dict:
    p = {"size": "m", "spacing": "normal", "font": "atkinson", "theme": "cream"}
    p.update({k: v for k, v in (prefs or {}).items() if v})
    return p


def _runs(paragraph: str, start_index: int, highlights: dict):
    """Yield (text, color_or_None) pieces and the next word index."""
    pieces, last, i = [], 0, start_index
    for m in WORD_RE.finditer(paragraph):
        if m.start() > last:
            pieces.append((paragraph[last:m.start()], None))
        pieces.append((m.group(), highlights.get(str(i))))
        last, i = m.end(), i + 1
    if last < len(paragraph):
        pieces.append((paragraph[last:], None))
    # merge neighbours with the same colour (incl. the space between two highlighted words)
    merged: list[list] = []
    for j, (text, color) in enumerate(pieces):
        if color is None and text.isspace() and merged and j + 1 < len(pieces) and merged[-1][1] and pieces[j + 1][1] == merged[-1][1]:
            merged[-1][0] += text
        elif merged and merged[-1][1] == color:
            merged[-1][0] += text
        else:
            merged.append([text, color])
    return merged, i


def safe_filename(title: str, ext: str) -> str:
    base = re.sub(r"[^A-Za-z0-9 _-]+", "", title or "reading").strip().replace(" ", "_")[:60] or "reading"
    return f"{base}_WordWise.{ext}"


# ---------- Word (.docx) ----------

def to_docx(doc: dict, highlights: dict, prefs: dict | None) -> bytes:
    p = _prefs(prefs)
    line, letter_em, _ = SPACING.get(p["spacing"], SPACING["normal"])
    size = SIZES_PT.get(p["size"], 14)
    bg, fg = THEMES.get(p["theme"], THEMES["cream"])

    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = FONTS.get(p["font"], "Verdana")
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), normal.font.name)
    normal.font.size = Pt(size)
    normal.font.color.rgb = RGBColor.from_string(fg.lstrip("#"))
    normal.paragraph_format.line_spacing = line
    normal.paragraph_format.space_after = Pt(size * 0.8)

    # Page colour (Word shows it in Print Layout; it is a display setting, not printed by default)
    background = OxmlElement("w:background")
    background.set(qn("w:color"), bg.lstrip("#"))
    d.element.insert(0, background)
    display = OxmlElement("w:displayBackgroundShape")
    d.settings.element.append(display)

    title = d.add_heading(doc.get("title", "Reading"), level=1)
    for r in title.runs:
        r.font.color.rgb = RGBColor.from_string("1B3366" if p["theme"] != "dark" else "FBF8F1")

    many = len(doc.get("sections", [])) > 1
    i = 0
    for s in doc.get("sections", []):
        if many:
            h = d.add_heading(s["label"], level=3)
            for r in h.runs:
                r.font.color.rgb = RGBColor.from_string("3C5A8F" if p["theme"] != "dark" else "C5CCDA")
        for para in s["paragraphs"]:
            runs, i = _runs(para, i, highlights or {})
            dp = d.add_paragraph()
            for text, color in runs:
                run = dp.add_run(text)
                if letter_em:
                    spacing = OxmlElement("w:spacing")
                    spacing.set(qn("w:val"), str(round(letter_em * size * 20)))  # twentieths of a point
                    run._r.get_or_add_rPr().append(spacing)
                if color in HL_DOCX:
                    run.font.highlight_color = HL_DOCX[color]
                    run.font.color.rgb = RGBColor.from_string("0F1F3D")

    foot = d.add_paragraph()
    foot.add_run("Exported from WordWise.").font.size = Pt(9)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---------- Web page (.html) ----------

def to_html(doc: dict, highlights: dict, prefs: dict | None) -> str:
    p = _prefs(prefs)
    line, letter_em, word_em = SPACING.get(p["spacing"], SPACING["normal"])
    bg, fg = THEMES.get(p["theme"], THEMES["cream"])
    heading = "#1B3366" if p["theme"] != "dark" else "#FBF8F1"

    body, i = [], 0
    many = len(doc.get("sections", [])) > 1
    for s in doc.get("sections", []):
        if many:
            body.append(f'<h2 class="label">{html.escape(s["label"])}</h2>')
        for para in s["paragraphs"]:
            runs, i = _runs(para, i, highlights or {})
            parts = [
                f'<mark class="hl-{c}">{html.escape(t)}</mark>' if c in HL_CSS else html.escape(t)
                for t, c in runs
            ]
            body.append("<p>" + "".join(parts) + "</p>")

    marks = "".join(f".hl-{k}{{background:{v};color:#0F1F3D;padding:0 2px;border-radius:3px}}" for k, v in HL_CSS.items())
    title = html.escape(doc.get("title", "Reading"))
    return f"""<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title} · WordWise</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Atkinson+Hyperlegible:wght@400;700&family=Lexend:wght@400;600&family=Libre+Baskerville:wght@700&display=swap">
<style>
body{{margin:0;background:{bg};color:{fg};font-family:{CSS_FONTS.get(p["font"])};font-size:{SIZES_PX.get(p["size"], 18)}px;
line-height:{line};letter-spacing:{letter_em}em;word-spacing:{word_em}em}}
main{{max-width:68ch;margin:0 auto;padding:48px 24px}}
h1{{font-family:'Libre Baskerville',Georgia,serif;color:{heading};letter-spacing:0;word-spacing:0;line-height:1.3}}
h2.label{{font-size:.8em;color:{heading};opacity:.75;margin:2em 0 .5em}}
p{{margin:0 0 1.1em}}
mark{{color:#0F1F3D}}{marks}
footer{{font-size:.75em;opacity:.7;margin-top:3em}}
@media print{{body{{-webkit-print-color-adjust:exact;print-color-adjust:exact}}}}
</style></head>
<body><main>
<h1>{title}</h1>
{chr(10).join(body)}
<footer>Exported from WordWise with your reading settings and highlights.</footer>
</main></body></html>"""
