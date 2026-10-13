"""Turn an uploaded file (or pasted text) into a clean reading document.

A document is a plain dict so it can live in a dcc.Store:
    {"title": str, "kind": "pdf" | "slides" | "docx" | "text",
     "sections": [{"label": "Page 1", "paragraphs": [str, ...]}, ...],
     "note": str | None}            # e.g. "We loaded the first 60 pages."
"""

from __future__ import annotations

import base64
import io
import json
import re
from pathlib import Path

from config import DATA_DIR

MAX_BYTES = 15 * 1024 * 1024  # 15 MB
MAX_SECTIONS = 60  # pages or slides
MAX_CHARS = 250_000  # roughly a 100-page book chapter set

SAMPLES_PATH = DATA_DIR / "sample" / "passages.json"


class DocumentError(ValueError):
    """A friendly, user-facing problem with the file."""


# ---------- cleanup helpers ----------

_SENTENCE_END = re.compile(r"[.!?:;\"'”’)\]]$")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # re-join words split across lines
    return text.strip()


def _paragraphs(text: str) -> list[str]:
    """Rebuild paragraphs from text that has a hard line break on every line (typical of PDFs)."""
    text = _clean(text)
    if not text:
        return []
    blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
    out: list[str] = []
    for block in blocks:
        lines = [ln.strip() for ln in block.split("\n") if ln.strip()]
        longest = max((len(ln) for ln in lines), default=0)
        current, last_line = "", ""
        for ln in lines:
            if current:
                # A short line that ends a sentence usually ends a paragraph.
                if _SENTENCE_END.search(last_line) and len(last_line) < 0.6 * longest and ln[:1].isupper():
                    out.append(current)
                    current = ln
                else:
                    current += " " + ln
            else:
                current = ln
            last_line = ln
        if current:
            out.append(current)
    return [re.sub(r"\s+", " ", p).strip() for p in out if len(p.strip()) > 1]


def _finish(title: str, kind: str, sections: list[dict], truncated: bool) -> dict:
    sections = [s for s in sections if s["paragraphs"]]
    if not sections:
        raise DocumentError(
            "We couldn't find any words in that file. If it is a scan or a photo of pages, "
            "try a PDF with selectable text."
        )
    total, kept = 0, []
    for s in sections:
        total += sum(len(p) for p in s["paragraphs"])
        kept.append(s)
        if total > MAX_CHARS:
            truncated = True
            break
    note = f"This is a long file, so we loaded the first {len(kept)} {'slides' if kind == 'slides' else 'pages'}." if truncated else None
    return {"title": title, "kind": kind, "sections": kept, "note": note}


# ---------- readers ----------

def _read_pdf(data: bytes, title: str) -> dict:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")
        pages = reader.pages
    except Exception as exc:  # noqa: BLE001 - any parse failure is a user-facing message
        raise DocumentError("We couldn't open that PDF. It may be damaged or password-protected.") from exc
    sections = [
        {"label": f"Page {i + 1}", "paragraphs": _paragraphs(page.extract_text() or "")}
        for i, page in enumerate(pages[:MAX_SECTIONS])
    ]
    return _finish(title, "pdf", sections, truncated=len(pages) > MAX_SECTIONS)


def _read_pptx(data: bytes, title: str) -> dict:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError("We couldn't open those slides. Try saving them as a PDF.") from exc
    sections = []
    for i, slide in enumerate(list(deck.slides)[:MAX_SECTIONS]):
        paras = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    line = "".join(r.text for r in p.runs).strip()
                    if line:
                        paras.append(line)
        sections.append({"label": f"Slide {i + 1}", "paragraphs": paras})
    return _finish(title, "slides", sections, truncated=len(deck.slides) > MAX_SECTIONS)


def _read_docx(data: bytes, title: str) -> dict:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise DocumentError("We couldn't open that Word file. Try saving it as a PDF.") from exc
    paras = [re.sub(r"\s+", " ", p.text).strip() for p in doc.paragraphs if p.text.strip()]
    return _finish(title, "docx", [{"label": "Document", "paragraphs": paras}], truncated=False)


def from_text(text: str, title: str = "Pasted text") -> dict:
    return _finish(title, "text", [{"label": "Text", "paragraphs": _paragraphs(text)}], truncated=False)


def from_upload(contents: str, filename: str) -> dict:
    """`contents` is the data URL from dcc.Upload."""
    try:
        _, b64 = contents.split(",", 1)
        data = base64.b64decode(b64)
    except Exception as exc:  # noqa: BLE001
        raise DocumentError("We couldn't read that file. Please try again.") from exc
    if len(data) > MAX_BYTES:
        raise DocumentError("That file is bigger than 15 MB. Try a smaller file or just the pages you need.")

    title = Path(filename or "Your reading").stem.replace("_", " ")
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        return _read_pdf(data, title)
    if ext == ".pptx":
        return _read_pptx(data, title)
    if ext == ".docx":
        return _read_docx(data, title)
    if ext in {".txt", ".md"}:
        return from_text(data.decode("utf-8", errors="replace"), title)
    raise DocumentError("We can read PDF, PowerPoint (.pptx), Word (.docx), and text (.txt) files.")


# ---------- samples ----------

def samples() -> list[dict]:
    return json.loads(SAMPLES_PATH.read_text(encoding="utf-8"))


def from_sample(sample_id: str) -> dict:
    for s in samples():
        if s["id"] == sample_id:
            return from_text("\n\n".join(s["paragraphs"]), s["title"])
    raise DocumentError("We couldn't find that sample.")


def section_text(doc: dict, index: int | None = None) -> str:
    """Text of one section, or the whole document when index is None."""
    sections = doc["sections"] if index is None else [doc["sections"][index]]
    return "\n\n".join(f"[{s['label']}]\n" + "\n".join(s["paragraphs"]) for s in sections)
